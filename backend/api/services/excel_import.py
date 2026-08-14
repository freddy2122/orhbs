"""Import Excel des fiches agents ORHS."""

from datetime import datetime
from io import BytesIO

from typing import Optional

from django.db import transaction
from openpyxl import Workbook, load_workbook

from api.constants.rhs_fields import EXCEL_AGENT_COLUMNS
from api.models import AgentSante, CampagneCollecte, ImportFichier, Structure
from api.services.declaration_sync import sync_declaration_from_agents


def generate_agent_template() -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Agents RHS"
    headers = [label for _, label in EXCEL_AGENT_COLUMNS]
    ws.append(headers)
    ws.append(
        [
            "MS-2024-00001",
            "Agossa",
            "Jean-Baptiste",
            "M",
            "1985-03-15",
            "Médecin spécialiste",
            "Médecin spécialiste",
            "Chirurgie générale",
            "Doctorat en Médecine",
            "UAC",
            "2010",
            "public",
            "permanent",
            "actif",
            "chu-mel",
            "Chirurgien",
            "2012-09-01",
            "",
            "2045-03-15",
            "+22990000000",
            "agent@example.bj",
        ]
    )
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _parse_date(value):
    if value is None or value == "":
        return None
    if hasattr(value, "date"):
        return value.date()
    if isinstance(value, datetime):
        return value.date()
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _parse_int(value):
    if value is None or value == "":
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _normalize_choice(value, mapping):
    if not value:
        return None
    key = str(value).strip().lower()
    return mapping.get(key, key)


SECTEUR_MAP = {
    "public": AgentSante.Secteur.PUBLIC,
    "publique": AgentSante.Secteur.PUBLIC,
    "prive": AgentSante.Secteur.PRIVE,
    "privé": AgentSante.Secteur.PRIVE,
    "confessionnel": AgentSante.Secteur.CONFESSIONNEL,
}

CONTRAT_MAP = {
    "permanent": AgentSante.TypeContrat.PERMANENT,
    "contractuel": AgentSante.TypeContrat.CONTRACTUEL,
    "prestataire": AgentSante.TypeContrat.PRESTATAIRE,
    "stagiaire": AgentSante.TypeContrat.STAGIAIRE,
}

STATUT_MAP = {
    "actif": AgentSante.StatutAgent.ACTIF,
    "détaché": AgentSante.StatutAgent.DETACHE,
    "detache": AgentSante.StatutAgent.DETACHE,
    "congé": AgentSante.StatutAgent.CONGE,
    "conge": AgentSante.StatutAgent.CONGE,
    "retraité": AgentSante.StatutAgent.RETRAITE,
    "retraite": AgentSante.StatutAgent.RETRAITE,
    "suspendu": AgentSante.StatutAgent.SUSPENDU,
}


def import_agents_excel(
    file_obj,
    campagne: CampagneCollecte,
    user,
    structure: Optional[Structure] = None,
) -> ImportFichier:
    log = ImportFichier.objects.create(
        campagne=campagne,
        structure=structure,
        nom_fichier=getattr(file_obj, "name", "import.xlsx"),
        importe_par=user,
    )

    try:
        wb = load_workbook(file_obj, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise ValueError("Fichier vide.")

        header = [str(c).strip().lower() if c else "" for c in rows[0]]
        col_keys = [key for key, label in EXCEL_AGENT_COLUMNS]
        col_index = {}
        for idx, cell in enumerate(header):
            for key, label in EXCEL_AGENT_COLUMNS:
                if cell == label.lower() or cell == key.lower():
                    col_index[key] = idx

        missing = [k for k in ("matricule", "nom", "prenom", "structure_code") if k not in col_index]
        if missing:
            raise ValueError(f"Colonnes obligatoires manquantes : {', '.join(missing)}")

        errors = []
        ok = 0
        total = 0
        structures_touched = set()

        with transaction.atomic():
            for line_no, row in enumerate(rows[1:], start=2):
                if not row or all(c is None or str(c).strip() == "" for c in row):
                    continue
                total += 1

                def cell(key):
                    idx = col_index.get(key)
                    if idx is None:
                        return None
                    return row[idx] if idx < len(row) else None

                matricule = str(cell("matricule") or "").strip()
                if not matricule:
                    errors.append({"ligne": line_no, "erreur": "Matricule obligatoire."})
                    continue

                structure_code = str(cell("structure_code") or "").strip().lower()
                try:
                    struct = Structure.objects.get(code=structure_code, actif=True)
                except Structure.DoesNotExist:
                    errors.append({"ligne": line_no, "erreur": f"Structure « {structure_code} » introuvable."})
                    continue

                if structure and struct.id != structure.id:
                    errors.append({"ligne": line_no, "erreur": "Structure hors périmètre autorisé."})
                    continue

                sexe_raw = str(cell("sexe") or "M").strip().upper()[:1]
                if sexe_raw not in ("M", "F"):
                    errors.append({"ligne": line_no, "erreur": "Sexe invalide (M ou F)."})
                    continue

                secteur_raw = _normalize_choice(cell("secteur"), SECTEUR_MAP)
                contrat_raw = _normalize_choice(cell("type_contrat"), CONTRAT_MAP)
                statut_raw = _normalize_choice(cell("statut_agent"), STATUT_MAP)

                defaults = {
                    "structure": struct,
                    "nom": str(cell("nom") or "").strip(),
                    "prenom": str(cell("prenom") or "").strip(),
                    "sexe": sexe_raw,
                    "date_naissance": _parse_date(cell("date_naissance")),
                    "profession": str(cell("profession") or "Autre").strip(),
                    "grade": str(cell("grade") or "").strip(),
                    "specialite": str(cell("specialite") or "").strip(),
                    "diplome_principal": str(cell("diplome_principal") or "").strip(),
                    "ecole_formation": str(cell("ecole_formation") or "").strip(),
                    "annee_diplome": _parse_int(cell("annee_diplome")),
                    "secteur": secteur_raw or AgentSante.Secteur.PUBLIC,
                    "type_contrat": contrat_raw or AgentSante.TypeContrat.PERMANENT,
                    "statut_agent": statut_raw or AgentSante.StatutAgent.ACTIF,
                    "poste_occupe": str(cell("poste_occupe") or "").strip(),
                    "date_prise_service": _parse_date(cell("date_prise_service")),
                    "date_fin_contrat": _parse_date(cell("date_fin_contrat")),
                    "depart_retraite_prevu": _parse_date(cell("depart_retraite_prevu")),
                    "telephone": str(cell("telephone") or "").strip(),
                    "email": str(cell("email") or "").strip(),
                    "source_fichier": log.nom_fichier,
                    "actif": True,
                }

                if not defaults["nom"]:
                    errors.append({"ligne": line_no, "erreur": "Nom obligatoire."})
                    continue

                AgentSante.objects.update_or_create(
                    campagne=campagne,
                    matricule=matricule,
                    defaults=defaults,
                )
                structures_touched.add(struct.id)
                ok += 1

            for struct_id in structures_touched:
                sync_declaration_from_agents(campagne, Structure.objects.get(pk=struct_id))

        log.lignes_total = total
        log.lignes_ok = ok
        log.lignes_erreur = len(errors)
        log.rapport_erreurs = errors[:100]
        log.statut = (
            ImportFichier.StatutImport.ERREUR
            if ok == 0 and errors
            else ImportFichier.StatutImport.TERMINE
        )
        log.save()
        return log

    except Exception as exc:
        log.statut = ImportFichier.StatutImport.ERREUR
        log.rapport_erreurs = [{"ligne": 0, "erreur": str(exc)}]
        log.save()
        return log
