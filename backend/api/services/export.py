import io
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from django.http import HttpResponse


def export_agents_excel(agents):
    """Exporte la liste des agents au format Excel."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Agents de santé"

    # Styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0F7B4F", end_color="0F7B4F", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # En-têtes
    headers = [
        "Matricule", "Nom", "Prénom", "Sexe", "Profession", "Grade", "Spécialité",
        "Structure", "Département", "Zone sanitaire", "Secteur", "Statut",
        "Type contrat", "Poste occupé", "Date prise service", "Date fin contrat",
        "Départ retraite prévu", "Téléphone", "Email", "Nationalité",
        "Diplôme principal", "École formation", "Année diplôme"
    ]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = border

    # Données
    for row, agent in enumerate(agents, 2):
        ws.cell(row=row, column=1, value=agent.matricule)
        ws.cell(row=row, column=2, value=agent.nom)
        ws.cell(row=row, column=3, value=agent.prenom)
        ws.cell(row=row, column=4, value=agent.get_sexe_display())
        ws.cell(row=row, column=5, value=agent.profession)
        ws.cell(row=row, column=6, value=agent.grade)
        ws.cell(row=row, column=7, value=agent.specialite)
        ws.cell(row=row, column=8, value=agent.structure.nom if agent.structure else "")
        ws.cell(row=row, column=9, value=agent.structure.departement.nom if agent.structure and agent.structure.departement else "")
        ws.cell(row=row, column=10, value=agent.structure.zone_sanitaire.nom if agent.structure and agent.structure.zone_sanitaire else "")
        ws.cell(row=row, column=11, value=agent.get_secteur_display())
        ws.cell(row=row, column=12, value=agent.get_statut_agent_display())
        ws.cell(row=row, column=13, value=agent.get_type_contrat_display())
        ws.cell(row=row, column=14, value=agent.poste_occupe)
        ws.cell(row=row, column=15, value=agent.date_prise_service.strftime("%d/%m/%Y") if agent.date_prise_service else "")
        ws.cell(row=row, column=16, value=agent.date_fin_contrat.strftime("%d/%m/%Y") if agent.date_fin_contrat else "")
        ws.cell(row=row, column=17, value=agent.depart_retraite_prevu.strftime("%d/%m/%Y") if agent.depart_retraite_prevu else "")
        ws.cell(row=row, column=18, value=agent.telephone)
        ws.cell(row=row, column=19, value=agent.email)
        ws.cell(row=row, column=20, value=agent.nationalite)
        ws.cell(row=row, column=21, value=agent.diplome_principal)
        ws.cell(row=row, column=22, value=agent.ecole_formation)
        ws.cell(row=row, column=23, value=agent.annee_diplome if agent.annee_diplome else "")

        # Appliquer les bordures
        for col in range(1, 24):
            ws.cell(row=row, column=col).border = border

    # Ajuster la largeur des colonnes
    column_widths = [15, 20, 20, 8, 20, 15, 20, 30, 20, 20, 12, 15, 15, 25, 15, 15, 18, 15, 25, 15, 30, 30, 12]
    for col, width in enumerate(column_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width

    # Sauvegarder dans un buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="agents_rhs_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx"'
    return response


def export_declarations_excel(declarations):
    """Exporte les déclarations RHS au format Excel."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Déclarations RHS"

    # Styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0F7B4F", end_color="0F7B4F", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # En-têtes
    headers = [
        "Structure", "Département", "Zone sanitaire", "Campagne", "Statut",
        "Effectif total", "Femmes", "Hommes", "Médecins", "Médecins généralistes",
        "Médecins spécialistes", "Infirmiers", "Infirmiers auxiliaires", "Sages-femmes",
        "Sages-femmes auxiliaires", "Pharmaciens", "Techniciens labo", "Techniciens imagerie",
        "Agents santé communautaire", "Chirurgiens dentistes", "Kinésithérapeutes",
        "Autres paramédicaux", "Administratifs", "Agents entretien", "Autre personnel",
        "Postes budgétés", "Postes pourvus", "Postes vacants",
        "Départs retraite 6 mois", "Départs retraite 12 mois", "Date soumission",
        "Date validation département", "Date validation nationale"
    ]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = border

    # Données
    for row, decl in enumerate(declarations, 2):
        ws.cell(row=row, column=1, value=decl.structure.nom if decl.structure else "")
        ws.cell(row=row, column=2, value=decl.structure.departement.nom if decl.structure and decl.structure.departement else "")
        ws.cell(row=row, column=3, value=decl.structure.zone_sanitaire.nom if decl.structure and decl.structure.zone_sanitaire else "")
        ws.cell(row=row, column=4, value=decl.campagne.libelle if decl.campagne else "")
        ws.cell(row=row, column=5, value=decl.get_statut_display())
        ws.cell(row=row, column=6, value=decl.effectif_total)
        ws.cell(row=row, column=7, value=decl.dont_femmes)
        ws.cell(row=row, column=8, value=decl.dont_hommes)
        ws.cell(row=row, column=9, value=decl.medecins)
        ws.cell(row=row, column=10, value=decl.medecins_generalistes)
        ws.cell(row=row, column=11, value=decl.medecins_specialistes)
        ws.cell(row=row, column=12, value=decl.infirmiers)
        ws.cell(row=row, column=13, value=decl.infirmiers_auxiliaires)
        ws.cell(row=row, column=14, value=decl.sages_femmes)
        ws.cell(row=row, column=15, value=decl.sages_femmes_auxiliaires)
        ws.cell(row=row, column=16, value=decl.pharmaciens)
        ws.cell(row=row, column=17, value=decl.techniciens_laboratoire)
        ws.cell(row=row, column=18, value=decl.techniciens_imagerie)
        ws.cell(row=row, column=19, value=decl.agents_sante_communautaire)
        ws.cell(row=row, column=20, value=decl.chirurgiens_dentistes)
        ws.cell(row=row, column=21, value=decl.kinesitherapeutes)
        ws.cell(row=row, column=22, value=decl.autres_paramedicaux)
        ws.cell(row=row, column=23, value=decl.administratifs)
        ws.cell(row=row, column=24, value=decl.agents_entretien)
        ws.cell(row=row, column=25, value=decl.autre_personnel)
        ws.cell(row=row, column=26, value=decl.postes_budgetes)
        ws.cell(row=row, column=27, value=decl.postes_pourvus)
        ws.cell(row=row, column=28, value=decl.postes_vacants)
        ws.cell(row=row, column=29, value=decl.departs_retraite_6_mois)
        ws.cell(row=row, column=30, value=decl.departs_retraite_12_mois)
        ws.cell(row=row, column=31, value=decl.date_soumission.strftime("%d/%m/%Y %H:%M") if decl.date_soumission else "")
        ws.cell(row=row, column=32, value=decl.date_valide_dept.strftime("%d/%m/%Y %H:%M") if decl.date_valide_dept else "")
        ws.cell(row=row, column=33, value=decl.date_valide_national.strftime("%d/%m/%Y %H:%M") if decl.date_valide_national else "")

        # Appliquer les bordures
        for col in range(1, 34):
            ws.cell(row=row, column=col).border = border

    # Ajuster la largeur des colonnes
    column_widths = [30, 20, 20, 25, 15, 12, 10, 10, 10, 15, 18, 12, 18, 15, 20, 12, 18, 18, 22, 20, 18, 20, 15, 18, 15, 12, 12, 12, 18, 18, 20, 22, 22]
    for col, width in enumerate(column_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width

    # Sauvegarder dans un buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="declarations_rhs_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx"'
    return response
